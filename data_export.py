"""Build user-owned MindPulse data exports in CSV, PDF, and DOCX formats."""
from __future__ import annotations

import csv
from datetime import datetime, timezone
from io import BytesIO, StringIO
from xml.sax.saxutils import escape


def _records(data):
    """Return ordered (section, list-of-records) pairs, omitting internal IDs."""
    sections = [
        ("Profile", [data.get("profile", {})] if data.get("profile") else []),
        ("Assessments", data.get("predictions", [])),
        ("Journal entries", data.get("journals", [])),
        ("Mood check-ins", data.get("checkins", [])),
        ("Habits", data.get("habits", [])),
        ("Goals", data.get("goals", [])),
    ]
    result = []
    for title, rows in sections:
        cleaned = []
        for row in rows:
            cleaned.append({k: v for k, v in dict(row).items() if k not in {"id", "user_id", "password_hash"}})
        result.append((title, cleaned))
    return result


def _csv_bytes(data):
    records = _records(data)
    fields = ["record_type"]
    for _, rows in records:
        for row in rows:
            for field in row:
                if field not in fields:
                    fields.append(field)
    output = StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=fields, extrasaction="ignore")
    writer.writeheader()
    for section, rows in records:
        for row in rows:
            safe_row = {"record_type": section}
            for key, value in row.items():
                if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
                    value = "'" + value
                safe_row[key] = value
            writer.writerow(safe_row)
    return output.getvalue().encode("utf-8-sig")


def _pdf_bytes(data):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=letter, rightMargin=0.65 * inch,
                            leftMargin=0.65 * inch, topMargin=0.65 * inch,
                            bottomMargin=0.65 * inch, title="MindPulse personal data export")
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="SmallValue", parent=styles["BodyText"], fontSize=8.5, leading=11, alignment=TA_LEFT, wordWrap="CJK"))
    story = [Paragraph("MindPulse — My Data", styles["Title"]),
             Paragraph("Personal data export. Generated " + datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"), styles["Italic"]),
             Spacer(1, 12),
             Paragraph("This file contains records associated with your MindPulse account. Wellbeing estimates are educational and are not medical diagnoses.", styles["BodyText"]),
             Spacer(1, 14)]
    for section, rows in _records(data):
        story.append(Paragraph(escape(section), styles["Heading2"]))
        if not rows:
            story.append(Paragraph("No records.", styles["BodyText"]))
            story.append(Spacer(1, 7))
            continue
        for index, row in enumerate(rows, 1):
            story.append(Paragraph(f"Record {index}", styles["Heading4"]))
            table_data = [[Paragraph("Field", styles["Heading5"]), Paragraph("Value", styles["Heading5"])]]
            for key, value in row.items():
                text = "" if value is None else str(value)
                table_data.append([Paragraph(escape(str(key).replace("_", " ").title()), styles["SmallValue"]), Paragraph(escape(text).replace("\n", "<br/>"), styles["SmallValue"])])
            table = Table(table_data, colWidths=[1.55 * inch, 5.45 * inch], repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef2ff")),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#d8deea")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.extend([table, Spacer(1, 8)])
    doc.build(story)
    return output.getvalue()


def _docx_bytes(data):
    from docx import Document
    from docx.shared import Inches, Pt

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    doc.add_heading("MindPulse — My Data", 0)
    doc.add_paragraph("Personal data export. Generated " + datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"))
    doc.add_paragraph("This file contains records associated with your MindPulse account. Wellbeing estimates are educational and are not medical diagnoses.")
    for title, rows in _records(data):
        doc.add_heading(title, level=1)
        if not rows:
            doc.add_paragraph("No records.")
            continue
        for index, row in enumerate(rows, 1):
            doc.add_heading(f"Record {index}", level=2)
            table = doc.add_table(rows=0, cols=2)
            table.style = "Light Shading Accent 1"
            for key, value in row.items():
                cells = table.add_row().cells
                cells[0].text = str(key).replace("_", " ").title()
                cells[1].text = "" if value is None else str(value)
                for paragraph in cells[0].paragraphs + cells[1].paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def build_export(data, file_format):
    """Return (bytes, mime_type, filename) for a supported personal-data export."""
    formatters = {
        "csv": (_csv_bytes, "text/csv; charset=utf-8", "mindpulse-my-data.csv"),
        "pdf": (_pdf_bytes, "application/pdf", "mindpulse-my-data.pdf"),
        "docx": (_docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "mindpulse-my-data.docx"),
    }
    if file_format not in formatters:
        raise ValueError("Unsupported export format")
    formatter, mimetype, filename = formatters[file_format]
    return formatter(data), mimetype, filename
